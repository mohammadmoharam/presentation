#!/usr/bin/env python3
"""Add the AI Dashboard screenshot appendix to AI-Executive-Overview.pptx.

Slides 1-3 (the executive story) keep their shapes and speaker notes; slide 2
only gains a small "View Dashboard Screenshots" button in the bottom-right
corner. Four supporting slides — one screenshot each — are appended after slide
3 and cross-linked with internal slide hyperlinks.

The script is idempotent: it removes any previously generated appendix slides
and link button before rebuilding them. Re-run it after update_slide_2.py,
which rebuilds slide 2 from scratch and therefore drops the link button.

Usage:
    pip install python-pptx pillow
    python3 add_screenshot_appendix.py
"""

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from generate_presentation import (
    LINK_SHAPE_NAME,
    OUT,
    add_screenshot_link,
    build_screenshot_appendix,
)

CORE_SLIDES = 3


def drop_slides_after(prs, keep):
    """Delete every slide after the first `keep` slides."""
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst)[keep:]:
        prs.part.drop_rel(sldId.rId)
        sldIdLst.remove(sldId)


def main():
    prs = Presentation(str(OUT))
    drop_slides_after(prs, CORE_SLIDES)

    slide_two = prs.slides[1]
    for shape in list(slide_two.shapes):
        if shape.name == LINK_SHAPE_NAME:
            shape._element.getparent().remove(shape._element)
    # Drop the stale hyperlink relationships left behind by the old button,
    # otherwise they keep the deleted appendix slide parts alive in the package.
    for rId, rel in list(slide_two.part.rels.items()):
        if rel.reltype == RT.SLIDE:
            slide_two.part.drop_rel(rId)

    appendix = build_screenshot_appendix(prs, slide_two)
    add_screenshot_link(slide_two, appendix[0])
    prs.save(str(OUT))
    print(f"Wrote {OUT} ({len(prs.slides)} slides, "
          f"{len(appendix)} screenshot appendix slides)")


if __name__ == "__main__":
    main()

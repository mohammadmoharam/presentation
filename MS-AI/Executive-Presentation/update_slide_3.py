#!/usr/bin/env python3
"""Rebuild slide 3 of AI-Executive-Overview.pptx in place.

Slides 1 and 2 are left untouched: only the shapes and speaker notes of the
third slide are cleared and regenerated with build_slide_three() from
generate_presentation.py, so the deck and the script stay in sync.

Usage:
    pip install python-pptx
    python3 update_slide_3.py
"""

from pptx import Presentation

from generate_presentation import OUT, build_slide_three


def main():
    prs = Presentation(str(OUT))
    slide = prs.slides[2]

    spTree = slide.shapes._spTree
    for shape in list(slide.shapes):
        spTree.remove(shape._element)
    slide.notes_slide.notes_text_frame.text = ""

    build_slide_three(slide)
    prs.save(str(OUT))
    print(f"Rebuilt slide 3 in {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()

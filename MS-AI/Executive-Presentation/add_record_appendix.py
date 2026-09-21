#!/usr/bin/env python3
"""Add the AI Priors record appendix and the tool/agent badges in place.

Slides 1-3 (the executive story) keep their shapes and speaker notes. Slide 1
gains a small "View AI Priors Records" button in the bottom-right corner and a
"Tool / Agent" badge in the footer; slide 2 gains the matching badge. Two
supporting slides — one screen recording each — are appended after the existing
screenshot appendix and cross-linked with internal slide hyperlinks.

The script is idempotent: previously generated record slides, link button and
badges are removed before being rebuilt. Run it after add_screenshot_appendix.py
(which deletes every slide after slide 3) and after update_slide_2.py (which
rebuilds slide 2 from scratch).

Usage:
    pip install python-pptx pillow
    python3 prepare_record_posters.py
    python3 add_record_appendix.py
"""

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from generate_presentation import (
    BADGE_SHAPE_NAME,
    LINK_SHAPE_NAME,
    OUT,
    RECORD_LINK_SHAPE_NAME,
    RECORD_SLIDE_MARKER,
    add_record_link,
    add_tool_agent_badge,
    build_record_appendix,
)


def is_record_slide(slide):
    return any(shape.name == RECORD_SLIDE_MARKER for shape in slide.shapes)


def drop_record_slides(prs):
    sldIdLst = prs.slides._sldIdLst
    for sldId, slide in zip(list(sldIdLst), list(prs.slides)):
        if is_record_slide(slide):
            prs.part.drop_rel(sldId.rId)
            sldIdLst.remove(sldId)


R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

SLIDE_ONE_NOTE = (
    "How it was built (only if asked): the capability runs on the Matcha AI "
    "Platform, with Claude Sonnet as the agent used to build it. Two optional "
    "demo recordings - the old manual flow and the new AI flow - are linked "
    "from the \"View AI Priors Records\" button in the bottom-right corner as "
    "supporting reference material."
)

SLIDE_TWO_NOTE = (
    "How it was built (only if asked): the dashboard uses the Matcha AI "
    "Platform, with Gemini Flash Lite 3.1 as the agent used to build it."
)


def append_note(slide, text):
    """Append the tool/agent sentence to the speaker notes exactly once."""
    tf = slide.notes_slide.notes_text_frame
    if "How it was built (only if asked)" in tf.text:
        return
    tf.text = f"{tf.text.rstrip()}\n\n{text}"


def drop_shapes(slide, names):
    """Remove the named shapes and any slide hyperlink they left behind.

    Only the relationships that no remaining shape references are dropped, so
    the existing slide-2 screenshot link keeps working.
    """
    for shape in list(slide.shapes):
        if shape.name in names:
            shape._element.getparent().remove(shape._element)

    in_use = {link.get(R_NS + "id")
              for link in slide.shapes._spTree.iter()
              if link.tag.endswith("}hlinkClick")}
    for rId, rel in list(slide.part.rels.items()):
        if rel.reltype == RT.SLIDE and rId not in in_use:
            slide.part.drop_rel(rId)


def main():
    prs = Presentation(str(OUT))
    drop_record_slides(prs)

    slide_one, slide_two = prs.slides[0], prs.slides[1]
    drop_shapes(slide_one, {RECORD_LINK_SHAPE_NAME, BADGE_SHAPE_NAME})
    drop_shapes(slide_two, {BADGE_SHAPE_NAME})

    records = build_record_appendix(prs, slide_one)
    add_record_link(slide_one, records[0])
    add_tool_agent_badge(slide_one, "Claude Sonnet")
    add_tool_agent_badge(slide_two, "Gemini Flash Lite 3.1")

    append_note(slide_one, SLIDE_ONE_NOTE)
    append_note(slide_two, SLIDE_TWO_NOTE)

    kept = [sh.name for sh in slide_two.shapes if sh.name == LINK_SHAPE_NAME]
    prs.save(str(OUT))
    print(f"Wrote {OUT} ({len(prs.slides)} slides, "
          f"{len(records)} record appendix slides, "
          f"slide-2 screenshot link kept: {bool(kept)})")


if __name__ == "__main__":
    main()
